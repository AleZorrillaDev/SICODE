import os
import cv2
import fitz  # PyMuPDF
import numpy as np
import pandas as pd
from pdf2image import convert_from_path
import shutil
import unicodedata
import openpyxl
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.table import Table, TableStyleInfo
from typing import List, Dict, Any, Optional, Callable

# TODO: Make these paths configurable via settings or environment variables
POPPLER_PATH = r"C:\poppler\Library\bin"
DEFAULT_SIGNATURES_DIR = r"D:\FIRMAS_AUXILIARES_COACTIVOS"

def quitar_tildes(texto: str) -> str:
    """Normaliza texto: quita tildes y acentos."""
    return ''.join(
        c for c in unicodedata.normalize('NFD', texto)
        if unicodedata.category(c) != 'Mn'
    )

def extraer_zonas_de_firma(pdf_path: str, salida_path: str, divisiones: int = 3):
    """Extrae la última página del PDF y la divide en zonas para buscar firmas."""
    # Convert only the last page to image
    # We use poppler_path if it exists, otherwise hope it's in system PATH
    kwargs = {}
    if os.path.exists(POPPLER_PATH):
        kwargs['poppler_path'] = POPPLER_PATH
        
    # Let's use PyMuPDF to get total pages first
    try:
        with fitz.open(pdf_path) as doc:
            total_pages = doc.page_count
        
        paginas = convert_from_path(pdf_path, first_page=total_pages, last_page=total_pages, **kwargs)
    except Exception as e:
        print(f"Error convirtiendo PDF {pdf_path} a imagen: {e}")
        return [], []
    
    if not paginas:
        return [], []
        
    ultima_pagina = paginas[0]
    ancho, alto = ultima_pagina.size
    nombre_base = os.path.splitext(os.path.basename(pdf_path))[0]
    zona_alto = alto // divisiones
    zonas_extraidas = []

    for i in range(divisiones):
        y_inicio = i * zona_alto
        y_fin = (i + 1) * zona_alto if i < divisiones - 1 else alto
        zona = ultima_pagina.crop((0, y_inicio, ancho, y_fin))
        ruta_salida = os.path.join(salida_path, f"{nombre_base}_zona_{i+1}.png")
        zona.save(ruta_salida)
        zonas_extraidas.append(ruta_salida)

    return zonas_extraidas, paginas

def comparar_firmas(zonas: List[str], carpeta_firmas: str, umbral: float = 0.7) -> Dict[str, str]:
    """Compara las zonas extraídas con las firmas de referencia en la carpeta."""
    if not os.path.isdir(carpeta_firmas):
        return {"Firmado": "Error", "Firmante": "Carpeta de firmas no encontrada"}

    firmas = [f for f in os.listdir(carpeta_firmas) if f.lower().endswith((".png", ".jpg", ".jpeg"))]
    coincidencias = []

    for firma_archivo in firmas:
        nombre_firma = os.path.splitext(firma_archivo)[0].split("_")[0].strip()
        firma_path = os.path.join(carpeta_firmas, firma_archivo)
        
        # Read reference signature
        firma_ref = cv2.imread(firma_path, cv2.IMREAD_GRAYSCALE)
        if firma_ref is None:
            continue

        for zona_path in zonas:
            zona = cv2.imread(zona_path, cv2.IMREAD_GRAYSCALE)
            if zona is None:
                continue
                
            # Skip if zone is smaller than reference
            if zona.shape[0] < firma_ref.shape[0] or zona.shape[1] < firma_ref.shape[1]:
                continue

            resultado = cv2.matchTemplate(zona, firma_ref, cv2.TM_CCOEFF_NORMED)
            _, max_val, _, _ = cv2.minMaxLoc(resultado)

            if max_val >= umbral:
                coincidencias.append((nombre_firma, max_val))
                break

    if coincidencias:
        mejor = max(coincidencias, key=lambda x: x[1])
        return {"Firmado": "Sí", "Firmante": mejor[0]}
    else:
        return {"Firmado": "No", "Firmante": ""}

def buscar_auxiliar_en_pdf(pdf_path: str, ruta_txt_auxiliares: str) -> str:
    """Busca nombres de auxiliares en el texto del PDF."""
    if not os.path.exists(ruta_txt_auxiliares):
        return "LISTA NO ENCONTRADA"

    with open(ruta_txt_auxiliares, "r", encoding="utf-8") as f:
        auxiliares = [quitar_tildes(line.strip().upper()) for line in f if line.strip()]

    texto_completo = ""
    try:
        with fitz.open(pdf_path) as doc:
            for page in doc:
                texto_completo += page.get_text().upper()
        texto_completo = quitar_tildes(texto_completo)
    except Exception as e:
        print(f"Error extrayendo texto de {pdf_path}: {e}")
        return "ERROR LECTURA"

    mejor_match = ""
    max_palabras = 0

    for nombre in auxiliares:
        palabras = nombre.split()
        # Cuenta cuántas palabras del nombre están en el texto del PDF
        coincidencias = sum(1 for palabra in palabras if palabra in texto_completo)
        if coincidencias > max_palabras and coincidencias >= 2:  # al menos 2 palabras
            mejor_match = nombre
            max_palabras = coincidencias

    return mejor_match if mejor_match else "SIN CONSIGNA"

def procesar_carpeta(
    carpeta_pdfs: str, 
    carpeta_firmas: str = DEFAULT_SIGNATURES_DIR, 
    on_progress: Optional[Callable[[int, int, str], None]] = None
) -> Dict[str, Any]:
    """Procesa todos los PDFs en una carpeta y genera un reporte."""
    
    carpeta_temp = os.path.join(carpeta_pdfs, "__temp_zonas__")
    os.makedirs(carpeta_temp, exist_ok=True)
    
    resultados = []
    archivos_pdf = [f for f in os.listdir(carpeta_pdfs) if f.lower().endswith(".pdf")]
    total = len(archivos_pdf)
    
    ruta_txt_auxiliares = os.path.join(carpeta_firmas, "AUXILIARESCOACTIVOS.txt")

    for i, archivo in enumerate(archivos_pdf):
        ruta_pdf = os.path.join(carpeta_pdfs, archivo)
        nombre_pdf = archivo
        
        if on_progress:
            on_progress(i + 1, total, nombre_pdf)

        try:
            # 1. Extraer zonas de firma
            zonas, _ = extraer_zonas_de_firma(ruta_pdf, carpeta_temp)
            
            # 2. Comparar con firmas de referencia
            res_firmas = comparar_firmas(zonas, carpeta_firmas)
            
            # 3. Buscar auxiliar en texto
            auxiliar = buscar_auxiliar_en_pdf(ruta_pdf, ruta_txt_auxiliares)
            
            # 4. Consolidar resultado (Mapping to the keys used in router)
            resultados.append({
                "filename": nombre_pdf,
                "firmante": res_firmas["Firmante"],
                "status": "Firma Válida" if res_firmas["Firmado"] == "Sí" else "Sin Firma",
                "valid": res_firmas["Firmado"] == "Sí",
                "auxiliar": auxiliar,
                "expediente": "---", 
                "ruc": "---"
            })
            
        except Exception as e:
            print(f"⚠️ Error en {nombre_pdf}: {e}")
            resultados.append({
                "filename": nombre_pdf,
                "firmante": "ERROR",
                "status": "Error",
                "valid": False,
                "auxiliar": "ERROR",
                "expediente": "---",
                "ruc": "---"
            })

    # Cleanup temp zones
    if os.path.exists(carpeta_temp):
        shutil.rmtree(carpeta_temp)

    # Generate Excel Report
    reporte_filename = "REPORTE_SIGNA_FIRMAS.xlsx"
    ruta_excel = os.path.join(carpeta_pdfs, reporte_filename)
    df_data = []
    for r in resultados:
        df_data.append({
            "Archivo": r["filename"],
            "Firmado": "SÍ" if r["valid"] else "NO",
            "Firmante": r["firmante"],
            "Auxiliar Coactivo": r["auxiliar"]
        })
    
    df = pd.DataFrame(df_data)
    
    # Save formatted excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "REPORTE"

    for row in dataframe_to_rows(df, index=False, header=True):
        ws.append(row)

    if not df.empty:
        n_filas = df.shape[0] + 1
        n_cols = df.shape[1]
        ultima_col = openpyxl.utils.get_column_letter(n_cols)
        rango_tabla = f"A1:{ultima_col}{n_filas}"

        tabla = Table(displayName="TablaSigna", ref=rango_tabla)
        estilo = TableStyleInfo(
            name="TableStyleMedium9", showFirstColumn=False,
            showLastColumn=False, showRowStripes=True, showColumnStripes=False
        )
        tabla.tableStyleInfo = estilo
        ws.add_table(tabla)

        # Auto-adjust columns
        for col in ws.columns:
            max_length = 0
            col_letter = col[0].column_letter
            for cell in col:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = max_length + 2

    wb.save(ruta_excel)

    return {
        "status": "success",
        "data": resultados,
        "reporte_excel": reporte_filename, # Return just the filename, router will handle serving it
        "summary": {
            "total": total,
            "valid": len([r for r in resultados if r["valid"]]),
            "error": len([r for r in resultados if not r["valid"]])
        }
    }
